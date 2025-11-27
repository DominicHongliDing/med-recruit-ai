import os
import google.generativeai as genai
import PyPDF2
import docx
import json
import smtplib
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart

def configure_ai(api_key: str):
    if api_key: genai.configure(api_key=api_key)

def extract_text_from_file(uploaded_file) -> str:
    try:
        if uploaded_file.name.endswith('.pdf'):
            reader = PyPDF2.PdfReader(uploaded_file)
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
            return text
        elif uploaded_file.name.endswith('.docx'):
            doc = docx.Document(uploaded_file)
            text = "\n".join(para.text for para in doc.paragraphs)
            return text
        elif uploaded_file.name.endswith('.txt'):
            return str(uploaded_file.read(), "utf-8")
        return ""
    except Exception as e:
        return f"Error reading file: {str(e)}"

# --- 多智能体分析流程 (Multi-Agent Workflow) ---
def analyze_batch_candidate(resume_text: str, jd_text: str, must_haves: str, role_type: str) -> dict:
    
    model = genai.GenerativeModel('gemini-2.0-flash')

    # --- AGENT 1: 数据提取 & 语言判断 ---
    prompt_agent_1 = f"""
    ROLE: Data Extraction Specialist.
    TASK: Extract data from the CV.
    
    CRITICAL STEP - LANGUAGE DETECTION:
    - Determine the candidate's likely native language based on their name (Chinese characters?) and university/experience location.
    - Field `language_preference`: Set to "Chinese" if they seem to be Chinese (e.g. Wang Xiaoming, graduated from Fudan). Set to "English" if they seem International/Non-Chinese speaker.
    
    CV TEXT: {resume_text[:15000]}
    
    EXTRACT JSON:
    1. Name (Keep original)
    2. Email
    3. language_preference ("Chinese" or "English")
    4. Education (University Name + Degree).
    5. Total Years of Experience (Number)
    6. List of Hard Skills
    7. Top 3 Papers (If applicable)
    
    OUTPUT: JSON only.
    """
    try:
        raw_data_response = model.generate_content(prompt_agent_1, generation_config={"response_mime_type": "application/json"})
        extracted_facts = json.loads(raw_data_response.text)
    except:
        extracted_facts = {"language_preference": "English"} # Default

    # --- AGENT 2: 风险分析师 (输出中文) ---
    prompt_agent_2 = f"""
    角色: 高级招聘风控专家。
    任务: 对比候选人简历与职位描述(JD)，找出风险点。
    
    JD: {jd_text}
    核心要求: {must_haves}
    候选人事实: {json.dumps(extracted_facts)}
    
    指令:
    - 请用**中文**回答。
    - 严格审查：工作年限是否造假？跳槽是否频繁？核心技能是否缺失？
    - 如果没有明显风险，请回答“未发现明显风险”。
    
    输出: 请列出 3 个简短的风险点（Bullet points）。
    """
    try:
        critique_response = model.generate_content(prompt_agent_2)
        critique_text = critique_response.text
    except:
        critique_text = "分析失败"

    # --- AGENT 3: 最终决策者 (输出中文) ---
    if "PI" in role_type or "Postdoc" in role_type:
        json_struct = '"bibliometrics": {"h_index": "Val", "total_citations": "Val", "total_paper_count": "Val"}, "representative_papers": [{"title":"", "journal":"", "significance":""}],'
    elif "Research Assistant" in role_type:
        json_struct = '"technical_skills": ["Skill 1"], "lab_experience_years": "Val", "project_participation": ["Proj 1"],'
    else: 
        json_struct = '"core_competencies": ["Comp 1"], "software_tools": ["Tool 1"],'

    prompt_agent_3 = f"""
    角色: 招聘决策官。
    输入:
    1. JD: {jd_text}
    2. 事实: {json.dumps(extracted_facts)}
    3. 风险报告: {critique_text}
    
    任务: 生成最终评估 JSON。
    **重要**: `summary`, `strengths`, `gaps`, `critique_notes` 必须用**中文**撰写，方便中国 HR 阅读。
    
    OUTPUT SCHEMA (JSON):
    {{
        "name": "{extracted_facts.get('name', 'Candidate')}",
        "email": "{extracted_facts.get('email', '')}",
        "language_preference": "{extracted_facts.get('language_preference', 'English')}",
        "fit_score": 0-100,
        "summary": "候选人中文画像总结。",
        "critique_notes": "{critique_text.replace(chr(10), ' ')}", 
        {json_struct}
        "strengths": ["优势1", "优势2"],
        "gaps": ["劣势1"]
    }}
    """
    
    try:
        final_response = model.generate_content(prompt_agent_3, generation_config={"response_mime_type": "application/json"})
        data = json.loads(final_response.text)
        if isinstance(data, list): return data[0]
        return data
    except Exception as e:
        return {"name": "Error", "fit_score": 0, "summary": f"AI Error: {str(e)}"}

# --- 智能多语言邮件生成器 ---
def generate_recruitment_email(candidate_data: dict, sender_info: dict, role_type: str) -> str:
    model = genai.GenerativeModel('gemini-2.0-flash')
    
    # 获取语言偏好 (Agent 1 提取的)
    lang = candidate_data.get('language_preference', 'Chinese')
    
    # 不同的 Prompt 策略
    if lang == "Chinese":
        # 中文邮件 Prompt
        draft_prompt = f"""
        写信人: {sender_info['name']} ({sender_info['org']}).
        收信人: {candidate_data.get('name')}. 职位: {role_type}.
        重点: {candidate_data.get('research_focus_area', '背景')}.
        
        任务: 写一封**中文**招聘邀请邮件。
        要求:
        1. 语气专业、诚恳。
        2. 提到对方的一个具体亮点（如论文或技能）。
        3. 邀请下周进行 15 分钟电话沟通。
        4. 不要包含 [占位符] 或 评分。
        """
    else:
        # 英文邮件 Prompt
        draft_prompt = f"""
        Writer: {sender_info['name']} ({sender_info['org']}).
        Target: {candidate_data.get('name')}. Role: {role_type}.
        Focus: {candidate_data.get('research_focus_area', 'Background')}.
        
        Task: Write a recruitment email in **English**.
        Requirements:
        1. Professional and personalized tone.
        2. Mention a specific highlight (Paper or Skill).
        3. Ask for a 15-min call next week.
        4. No placeholders like [Date] or Score.
        """

    draft = model.generate_content(draft_prompt).text
    return draft

def send_real_email(sender_email, sender_password, recipient_email, subject, body):
    try:
        message = MIMEMultipart()
        message['From'] = sender_email
        message['To'] = recipient_email
        message['Subject'] = subject
        message.attach(MIMEText(body, 'plain'))
        session = smtplib.SMTP('smtp.gmail.com', 587) 
        session.starttls()
        session.login(sender_email, sender_password) 
        session.sendmail(sender_email, recipient_email, message.as_string())
        session.quit()
        return True, "发送成功 / Email Sent Successfully"
    except Exception as e:
        return False, f"发送失败 / Failed: {str(e)}"
